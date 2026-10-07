"""第八批 §8.7 / §8.8 的接口回归（只跑隔离库，禁止任何真实外部调用）。

跑法（仓库根）：
    & .\\ops\\iso_checks.ps1 -Db crm_iso_f -Port 8016 \
        -Suites check_biz_docs,check_quote_lifecycle,check_biz_doc_freeze

覆盖点与断言：
1. §8.7 按报价版本冻结：币种（USD）、计价单位（套/箱）、有效期、付款/交付/贸易条款、
   客户与联系人抬头都进快照，并且**真的出图**后用 openpyxl 回读得到（不只查 JSON）；
2. §8.7 小数数量与附加费/优惠：明细金额 = 数量×单价，小计 + 费用 + 优惠 = 合计，
   都与报价版本行逐项一致；
3. §8.7 旧版本一致性：改当前客户/SKU（单位、规格、名称）后，旧版本再生成仍与该版
   条款一致；已出的旧文件快照与校验值一个字不动；
4. §8.7 缺字段不伪造：把版本行的单位快照清空（模拟历史版本）后生成，表上必须是
   "待核实"，不得出现当前 SKU 的单位；
5. §8.8 模板变量校验：拼错变量存不下去（422 且指名变量）；引用订单字段的模板在
   报价单生成时被拦（422，来源不可用）；
6. §8.8 草稿：`allow_draft=true` 时落成"草稿（非正式对外文件）"、正文不带 `{{ }}`、
   不可作废；默认（不传 allow_draft）直接拒绝；
7. §8.5 授权：`has_doc_permission` 按单据类型判来源模块权限——只有订单权限的人
   拿不到报价单（含 admin 全通过的反面）；
8. §8.9 幂等（第 11 节）：同一把请求键重发只出一份并返回原来那份，换新键才出新版；
9. §8.10 原件（第 9/10 节）：生成即存档、重复下载字节一致、状态副本与原件分得开、
   原件失踪时拒绝下载并告警（明确要求才给"由历史快照重建"的副本）。
"""

import asyncio
import hashlib
import os
import sys
import time
import urllib.error
import urllib.request
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from urllib.parse import urlparse

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import select, text  # noqa: E402

from app.core.config import settings  # noqa: E402
from app.core.database import SessionLocal  # noqa: E402
from scripts.check_review_regressions import BASE, call, login  # noqa: E402

PREFIX = f"CHKFZ{int(time.time())}"
FAILURES: list[str] = []


def check(label, actual, expected):
    good = actual == expected
    print(f'  {"OK  " if good else "FAIL"} {label}: {actual!r}（期望 {expected!r}）')
    if not good:
        FAILURES.append(label)


def check_true(label, condition, detail=""):
    print(f'  {"OK  " if condition else "FAIL"} {label}{f"：{detail}" if detail else ""}')
    if not condition:
        FAILURES.append(label)


def request(method, path, *, token, body=None, expect=200):
    status, result = call(method, path, body=body, token=token)
    assert status == expect, f"{method} {path} -> HTTP {status} {result}"
    return result


async def cleanup():
    """按外键顺序清掉本轮夹具（前缀唯一，不会碰别人的数据）。"""
    async with SessionLocal() as s:
        cust = "(select id from customers where name like :p)"
        quote = f"(select id from quotes where customer_id in {cust})"
        order = f"(select id from sales_orders where customer_id in {cust})"
        for sql in (
            f"delete from biz_docs where customer_id in {cust}",
            "delete from biz_doc_templates where name like :p",
            f"delete from quote_charges where quote_version_id in (select id from quote_versions where quote_id in {quote})",
            f"delete from quote_items where quote_version_id in (select id from quote_versions where quote_id in {quote})",
            f"delete from quote_versions where quote_id in {quote}",
            f"delete from quotes where customer_id in {cust}",
            f"delete from sales_order_items where order_id in {order}",
            f"delete from sales_orders where customer_id in {cust}",
            f"delete from contacts where customer_id in {cust}",
            # §8.14 复审：夹具现在会插主数据确认与整版快照，它们**引用 skus**，
            # 必须在删 SKU 之前先删（否则外键挡住，整个清理断在这里、
            # 夹具留在库里被守门套件抓出来）
            "delete from sku_master_versions where sku_id in "
            "(select id from skus where sku_code like :sku)",
            "delete from sku_field_authorities where sku_id in "
            "(select id from skus where sku_code like :sku)",
            "delete from skus where sku_code like :sku",
            # 夹具还建了一个**产品**（报价单要挂在产品/SKU 上）：原来只删 SKU，
            # 产品行留在库里，最后被 check_fixture_residue 揪出来。
            "delete from products where name like :p",
            f"delete from customers where name like :p",
        ):
            await s.execute(text(sql), {"p": f"{PREFIX}%", "sku": f"{PREFIX}%"})
        await s.commit()


def _cells(sheet):
    return [
        str(sheet.cell(row=row, column=column).value)
        for row in range(1, sheet.max_row + 1)
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(row=row, column=column).value is not None
    ]


def _download_raw(doc_id, token, query=""):
    """下载**原始字节 + 响应头**（§8.10 的验收只能在字节上验）。"""
    url = f"{BASE}/biz-docs/{doc_id}/download"
    if query:
        url = f"{url}?{query}"
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
    try:
        with urllib.request.urlopen(req, timeout=30) as response:
            # 返回 headers **对象本身**而不是 dict()：HTTP 头大小写不敏感，
            # 服务端实际发的可能是小写，转成 dict 之后 `.get("X-Doc-...")` 就取不到了。
            return response.read(), response.headers
    except urllib.error.HTTPError as exc:
        raise AssertionError(
            f"下载失败：HTTP {exc.code} {exc.read()[:200]!r}"
        ) from exc


def _download_xlsx(doc_id, token):
    """走真实下载接口拿 Excel 回读——渲染器/layout 只在字节流上才验得到。"""
    payload, _headers = _download_raw(doc_id, token)
    assert payload[:2] == b"PK", f"报价单下载的不是 xlsx：{payload[:8]!r}"
    return load_workbook(BytesIO(payload)).active


def _new_template(token, doc_type, name, body):
    result = request(
        "POST",
        "/biz-doc-templates",
        token=token,
        body={"doc_type": doc_type, "name": name, "body": body},
    )
    return result["data"]["id"]


async def main():
    import app.main  # noqa: F401

    _ = app.main
    assert urlparse(BASE).hostname in {"127.0.0.1", "localhost", "::1"}
    db = urlparse(settings.database_url)
    assert db.hostname in {"127.0.0.1", "localhost", "::1"}
    assert "iso" in db.path.lower() or "check" in db.path.lower(), (
        f"本脚本只允许跑在一次隔离库上，实际 {db.path}"
    )
    assert settings.dingtalk_push_off and settings.wecom_push_off
    assert not settings.scheduler_enabled

    from app.modules.bizdoc.model import has_doc_permission
    from app.modules.customer.model import Contact, Customer
    from app.modules.product.model import Product, Sku
    from app.modules.quote.model import Quote, QuoteCharge, QuoteItem, QuoteVersion

    token = login("admin", "admin123")
    await cleanup()
    now = datetime.now(UTC)
    quote_id = version_id = None

    async with SessionLocal() as s:
        from app.modules.user.model import User

        admin = (
            await s.execute(select(User).where(User.username == "admin"))
        ).scalars().one()
        customer = Customer(
            name=f"{PREFIX}冻结客户", owner_id=admin.id, status="active", level="A"
        )
        s.add(customer)
        await s.flush()
        contact = Contact(customer_id=customer.id, name=f"{PREFIX}联系人")
        s.add(contact)
        await s.flush()
        product = Product(name=f"{PREFIX}产品")
        s.add(product)
        await s.flush()
        sku = Sku(
            product_id=product.id,
            sku_code=f"{PREFIX}-SKU",
            name=f"{PREFIX}法兰",
            specification="DN50",
            unit="套",
        )
        s.add(sku)
        await s.flush()
        # §8.14 复审后，生成**正式件**要求明细引用的那一版主数据快照里有
        # 名称/规格/单位（未确认只能出草稿）。本套件测的是正式件的冻结行为，
        # 所以夹具先把这三个字段确认掉并落一版整版快照（明细引用 version_no=1）。
        from app.modules.product.model import SkuFieldAuthority, SkuMasterVersion  # noqa: E402

        display_values = {
            "name": sku.name,
            "specification": sku.specification,
            "unit": sku.unit,
        }
        for field, value in display_values.items():
            s.add(
                SkuFieldAuthority(
                    sku_id=sku.id,
                    field_name=field,
                    confirmed_version=1,
                    confirmed_value=value,
                    status="confirmed",
                    source_verified=False,
                )
            )
        s.add(
            SkuMasterVersion(
                sku_id=sku.id,
                version_no=1,
                values=display_values,
                source_summary={},
                confirmed_at=datetime.now(UTC),
                note="套件夹具：确认印给客户的三个字段",
            )
        )
        await s.flush()
        quote = Quote(
            quote_no=f"{PREFIX}-Q",
            customer_id=customer.id,
            contact_id=contact.id,
            owner_id=admin.id,
            status="approved",
            valid_until=(now + timedelta(days=30)).date(),
            created_at=now,
            updated_at=now,
        )
        s.add(quote)
        await s.flush()
        version = QuoteVersion(
            quote_id=quote.id,
            version_no=1,
            currency="USD",
            payment_terms="30% 预付，余款见提单副本",
            delivery_terms="FOB 宁波",
            trade_terms="FOB",
            # 2026-10-07：抬头与有效期改为取**版本快照**（原来实时读客户资料 / 报价主单）。
            # 本套件的夹具是直插库的，所以要自己把这三列写上 —— 这也更贴近真实：
            # 正常走接口创建版本时，服务层会写这三列。
            customer_name_snapshot=customer.name,
            contact_name_snapshot=contact.name,
            valid_until_snapshot=quote.valid_until,
            subtotal_amount=Decimal("850.00"),
            charge_amount=Decimal("80.00"),
            discount_amount=Decimal("-30.00"),
            total_amount=Decimal("900.00"),
            created_at=now,
        )
        s.add(version)
        await s.flush()
        quote.current_version_id = version.id
        s.add(
            QuoteItem(
                quote_version_id=version.id,
                sku_id=sku.id,
                sku_code_snapshot=sku.sku_code,
                sku_name_snapshot=sku.name,
                spec_snapshot=sku.specification,
                # §8.14 复审：明细要记"引用的那一版主数据快照"。不填的话，
                # 生成对客文件会判成"没有可引用的已确认主数据版本"→ 落草稿，
                # 而本套件测的是**正式件**的冻结行为。
                master_version_no=1,
                quantity=Decimal("7.500"),  # 小数数量
                quoted_price=Decimal("113.3333"),
                unit_snapshot="套",
            )
        )
        s.add(
            QuoteCharge(
                quote_version_id=version.id,
                charge_type="logistics",
                description="运费",
                amount=Decimal("80.00"),
                is_discount=False,
                sort_no=1,
            )
        )
        s.add(
            QuoteCharge(
                quote_version_id=version.id,
                charge_type="discount",
                description="整单优惠",
                amount=Decimal("-30.00"),
                is_discount=True,
                sort_no=2,
            )
        )
        await s.commit()
        quote_id, version_id, customer_id, sku_id, contact_id = (
            quote.id,
            version.id,
            customer.id,
            sku.id,
            contact.id,
        )
        valid_until = quote.valid_until.isoformat()

    try:
        print("=== 1. §8.7 生成报价单：币种/单位/条款按版本冻结 ===")
        generated = request(
            "POST", "/biz-docs/quote", token=token, body={"quote_version_id": version_id}
        )["data"]
        doc_id = generated["id"]
        snapshot = request("GET", f"/biz-docs/{doc_id}", token=token)["data"]["input_snapshot"]
        frozen = snapshot["frozen"]
        check("状态是有效件", generated["status"], "active")
        check("币种取自报价版本", frozen["terms"]["currency"], "USD")
        check("付款条件取自报价版本", frozen["terms"]["payment_terms"], "30% 预付，余款见提单副本")
        check("交付条件取自报价版本", frozen["terms"]["delivery_terms"], "FOB 宁波")
        check("有效期取自报价主单", frozen["terms"]["valid_until"], valid_until)
        check("客户抬头已冻结", frozen["header"]["customer_name"], f"{PREFIX}冻结客户")
        check("联系人抬头已冻结", frozen["header"]["contact_name"], f"{PREFIX}联系人")
        check("明细单位取版本快照", snapshot["items"][0]["unit"], "套")
        check("小数数量如实", str(snapshot["items"][0]["quantity"]), "7.5")
        check_true("没有未留存项", frozen["gaps"] == [], f"{frozen['gaps']}")

        # 金额口径：明细 = 数量×单价；小计 + 费用 + 优惠 = 合计
        item_amount = Decimal(str(snapshot["items"][0]["amount"]))
        check_true(
            "明细金额 = 数量 × 单价（版本快照）",
            item_amount == Decimal("7.5") * Decimal("113.3333"),
            f"{item_amount}",
        )
        charge_sum = sum(Decimal(str(row["amount"])) for row in snapshot["charges"])
        check(
            "小计 + 费用 + 优惠 = 合计",
            Decimal(str(snapshot["subtotal_amount"])) + charge_sum,
            Decimal(str(snapshot["total_amount"])),
        )

        print("=== 2. §8.7 真的出图：Excel 上能读到币种/单位/条款 ===")
        sheet = _download_xlsx(doc_id, token)
        texts = _cells(sheet)
        for expected in ("币种", "USD", "数量", "单位", "套", "有效期至", "付款条件",
                         "交付条件", "贸易条款", "FOB 宁波", "小计", "运费", "整单优惠", "合计"):
            check_true(f"Excel 上有「{expected}」", expected in texts, f"{texts}")
        unit_index = texts.index("单位")
        check("数量与单位同列组相邻", texts[unit_index - 1], "数量")
        check_true("Excel 上没有残留模板语法", not any("{{" in t for t in texts))

        print("=== 3. §8.7 改当前客户/SKU 后旧版本仍与该版一致 ===")
        async with SessionLocal() as s:
            customer = await s.get(Customer, customer_id)
            customer.name = f"{PREFIX}冻结客户-已改名"
            contact = await s.get(Contact, contact_id)
            contact.name = f"{PREFIX}联系人-已改名"
            sku = await s.get(Sku, sku_id)
            sku.name = f"{PREFIX}法兰-改名"
            sku.specification = "DN80"
            sku.unit = "箱"
            await s.commit()
        regenerated = request(
            "POST", "/biz-docs/quote", token=token, body={"quote_version_id": version_id}
        )["data"]
        check("重新生成是新增一版", regenerated["version"], 2)
        check("新版指向旧版", regenerated["parent_id"], doc_id)
        frozen_v2 = request("GET", f"/biz-docs/{regenerated['id']}", token=token)["data"][
            "input_snapshot"
        ]["frozen"]
        for field, expected in (
            ("currency", "USD"),
            ("payment_terms", "30% 预付，余款见提单副本"),
            ("delivery_terms", "FOB 宁波"),
            ("trade_terms", "FOB"),
            ("valid_until", valid_until),
        ):
            check(f"旧版本再生成后条款「{field}」不变", frozen_v2["terms"][field], expected)
        check("单位仍取版本快照（不是当前箱）", frozen_v2["items"][0]["unit"], "套")
        check("名称仍取版本快照", frozen_v2["items"][0]["name"], f"{PREFIX}法兰")
        check("规格仍取版本快照", frozen_v2["items"][0]["spec"], "DN50")
        old = request("GET", f"/biz-docs/{doc_id}", token=token)["data"]
        check("旧版快照一个字没动", old["input_snapshot"]["frozen"]["items"][0]["unit"], "套")
        check(
            "旧版校验值没变",
            old["content_sha256"],
            generated["content_sha256"],
        )

        print("=== 4. §8.7 历史版本缺单位：写「待核实」，不拿当前 SKU 单位 ===")
        async with SessionLocal() as s:
            row = (
                await s.execute(
                    select(QuoteItem).where(QuoteItem.quote_version_id == version_id)
                )
            ).scalars().one()
            row.unit_snapshot = None  # 模拟早于本修复的版本行
            await s.commit()
        legacy = request(
            "POST", "/biz-docs/quote", token=token, body={"quote_version_id": version_id}
        )["data"]
        legacy_snapshot = request("GET", f"/biz-docs/{legacy['id']}", token=token)["data"][
            "input_snapshot"
        ]
        check("缺单位时写待核实", legacy_snapshot["items"][0]["unit"], "待核实")
        check_true(
            "gap 里记了这条未留存",
            any(gap["field"].endswith(".unit") for gap in legacy_snapshot["frozen"]["gaps"]),
            f"{legacy_snapshot['frozen']['gaps']}",
        )
        legacy_texts = _cells(_download_xlsx(legacy["id"], token))
        check_true("Excel 上写「待核实」", "待核实" in legacy_texts, f"{legacy_texts}")
        check_true("Excel 上没有当前 SKU 的「箱」", "箱" not in legacy_texts, f"{legacy_texts}")

        print("=== 5. §8.8 模板变量校验：拼错/来源不支持都被拦 ===")
        status, result = call(
            "POST",
            "/biz-doc-templates",
            token=token,
            body={
                "doc_type": "quote_sheet",
                "name": f"{PREFIX}拼错模板",
                "body": "客户：{{customer.nmae}}",
            },
        )
        check("拼错变量存不下去（422）", status, 422)
        check_true(
            "报错指出了正确写法",
            "customer.name" in str(result),
            f"{result}",
        )
        preview = request(
            "POST",
            "/biz-doc-templates/preview",
            token=token,
            body={"doc_type": "order_sheet", "body": "付款：{{order.payment_terms}}"},
        )["data"]
        check("预览能解出变量", preview["status"], "ok")
        preview_bad = request(
            "POST",
            "/biz-doc-templates/preview",
            token=token,
            body={"doc_type": "order_sheet", "body": "付款：{{order.payment_termz}}"},
        )["data"]
        check("预览指出拼错", preview_bad["status"], "has_unresolved")
        check_true(
            "预览正文不带模板语法", "{{" not in preview_bad["body"], f"{preview_bad['body']}"
        )
        check_true(
            "预览给出可用变量清单",
            any("order.payment_terms" in item for item in preview_bad["analysis"]["available"]),
        )

        bad_template = _new_template(
            token, "quote_sheet", f"{PREFIX}引用订单字段", "付款：{{order.payment_terms}}"
        )
        status, result = call(
            "POST",
            "/biz-docs/quote",
            token=token,
            body={"quote_version_id": version_id, "template_id": bad_template},
        )
        check("未解析变量阻止出正式件（422）", status, 422)
        check_true(
            "报错指名了变量与原因",
            "order.payment_terms" in str(result) and "来源不可用" in str(result),
            f"{result}",
        )

        print("=== 6. §8.8 允许草稿：标草稿、正文无模板语法、不可作废 ===")
        async with SessionLocal() as s:
            from app.modules.order.model import SalesOrder, SalesOrderItem

            order = SalesOrder(
                order_no=f"{PREFIX}-SO",
                customer_id=customer_id,
                owner_id=admin.id,
                sales_owner_id=admin.id,
                total_amount=Decimal("100"),
                currency="CNY",
                status="pending",
                created_by=admin.id,
            )
            s.add(order)
            await s.flush()
            s.add(
                SalesOrderItem(
                    order_id=order.id,
                    sku_id=sku_id,
                    sku_snapshot=f"{PREFIX}法兰",
                    specification="DN80",
                    quantity=Decimal("2"),
                    unit_price=Decimal("50"),
                    amount=Decimal("100"),
                )
            )
            await s.commit()
            order_id = order.id

        order_template = _new_template(
            token, "order_sheet", f"{PREFIX}订单模板", "付款：{{order.payment_terms}}"
        )
        status, result = call(
            "POST",
            "/biz-docs/order",
            token=token,
            body={"order_id": order_id, "template_id": order_template},
        )
        check("默认不允许草稿出图（422）", status, 422)
        draft = request(
            "POST",
            "/biz-docs/order",
            token=token,
            body={"order_id": order_id, "template_id": order_template, "allow_draft": True},
        )["data"]
        check("明确允许时落草稿", draft["status"], "draft")
        check("草稿有明确标签", draft["status_label"], "草稿（非正式对外文件）")
        check_true("草稿记了缺哪些变量", bool(draft["token_issues"]), f"{draft['token_issues']}")
        draft_doc = request("GET", f"/biz-docs/{draft['id']}", token=token)["data"]
        check_true(
            "草稿正文不带模板语法",
            "{{" not in draft_doc["input_snapshot"]["body"],
            f"{draft_doc['input_snapshot']['body']}",
        )
        check_true(
            "草稿正文写明未解析",
            "未解析" in draft_doc["input_snapshot"]["body"]
            or "来源不可用" in draft_doc["input_snapshot"]["body"],
            f"{draft_doc['input_snapshot']['body']}",
        )
        status, _ = call(
            "POST", f"/biz-docs/{draft['id']}/void", token=token, body={"reason": "不该能作废草稿"}
        )
        check("草稿不可作废（422）", status, 422)

        print("=== 7. §8.5 单据授权按类型判来源模块权限 ===")

        class _User:
            def __init__(self, permissions, roles=()):
                self.permissions = set(permissions)
                self.roles = list(roles)

            def has(self, code):
                return code in self.permissions

        order_only = _User({"order:view", "order:manage"})
        quote_only = _User({"quote:view", "quote:manage"})
        check("只有订单权限：看得到下单文件", has_doc_permission(order_only, "order_sheet"), True)
        check("只有订单权限：拿不到报价单", has_doc_permission(order_only, "quote_sheet"), False)
        check("只有订单权限：拿不到打样单", has_doc_permission(order_only, "sample_request"), False)
        check("只有报价权限：看得到报价单", has_doc_permission(quote_only, "quote_sheet"), True)
        check(
            "只有报价权限：作废不了订单文件",
            has_doc_permission(quote_only, "order_sheet", manage=True),
            False,
        )
        check(
            "只有报价权限：能作废自己的报价单",
            has_doc_permission(quote_only, "quote_sheet", manage=True),
            True,
        )
        check(
            "管理员全通过",
            has_doc_permission(_User(set(), roles=["admin"]), "quote_sheet"),
            True,
        )

        print("=== 8. 列表按类型过滤，不把 order:view 当总开关 ===")
        listing = request("GET", "/biz-docs?limit=300", token=token)["data"]
        doc_types = {row["doc_type"] for row in listing if row["id"] in {doc_id, draft["id"]}}
        check("列表里能看到报价单", "quote_sheet" in doc_types, True)
        check("列表里能看到下单文件", "order_sheet" in doc_types, True)
        picked = request(
            "GET", f"/biz-docs?doc_type=quote_sheet&quote_id={quote_id}", token=token
        )["data"]
        check_true(
            "按报价单过滤有效",
            picked and all(row["doc_type"] == "quote_sheet" for row in picked),
            f"{picked}",
        )

        print("=== 9. §8.10 下载读存档原件：重复下载字节一致 ===")
        first_bytes, first_headers = _download_raw(doc_id, token)
        second_bytes, second_headers = _download_raw(doc_id, token)
        check_true(
            "两次下载的字节完全一致",
            first_bytes == second_bytes,
            f"{len(first_bytes)} 字节",
        )
        check(
            "响应头说明内容来自存档原件",
            first_headers.get("X-Doc-Content-Source"),
            "archived",
        )
        check("响应头说明有存档原件", first_headers.get("X-Doc-Archive-Status"), "archived")
        check(
            "下载字节的 SHA-256 等于库里登记的存档哈希",
            hashlib.sha256(first_bytes).hexdigest(),
            generated["archive"]["file_sha256"],
        )
        check_true(
            "字节哈希与内容哈希是两回事（不是同一个值）",
            generated["archive"]["file_sha256"] != generated["content_sha256"],
        )
        check(
            "存档原件记录了当时的渲染器版本",
            generated["archive"]["renderer_version"],
            "bizdoc-xlsx/1",
        )
        # 作废展示要另渲染：状态副本与存档原件**必须分得开**
        state_bytes, state_headers = _download_raw(doc_id, token, query="mode=state")
        check("状态副本标明自己不是原件", state_headers.get("X-Doc-Content-Source"), "state-copy")
        check_true("状态副本与存档原件不是同一份字节", state_bytes != first_bytes)
        state_texts = _cells(load_workbook(BytesIO(state_bytes)).active)
        check_true(
            "状态副本纸面上写明是副本",
            any("状态副本" in text for text in state_texts),
            f"{state_texts}",
        )
        status, _ = call("GET", f"/biz-docs/{doc_id}/download?mode=bogus", token=token)
        check("非法 mode 被拒（422）", status, 422)

        print("=== 10. §8.10 原件失踪：拒绝下载并告警，不静默用当前资料顶替 ===")
        from app.modules.bizdoc.model import BizDoc
        from app.modules.file import storage
        from app.modules.file.model import FileRecord

        async with SessionLocal() as s:
            row = await s.get(BizDoc, doc_id)
            record = await s.get(FileRecord, row.file_id)
            archived_path = storage.absolute_path(record.object_key)
            backup = archived_path.read_bytes()
        try:
            archived_path.unlink()
            status, result = call("GET", f"/biz-docs/{doc_id}/download", token=token)
            check("原件失踪时拒绝下载（410）", status, 410)
            check_true(
                "拒绝时说明了原因，且不假装给了原件",
                "存档原件不可用" in str(result) and "allow_rebuild" in str(result),
                f"{result}",
            )
            # "可告警"不能只是嘴上说说：拒绝这件事本身要留下审计痕迹，
            # 否则事后没人知道有原件失踪过（也不会有任何东西提醒运维去查）。
            from app.core.audit import AuditLog

            async with SessionLocal() as s:
                blocked = (
                    await s.execute(
                        select(AuditLog).where(
                            AuditLog.business_type == "biz_doc",
                            AuditLog.business_id == doc_id,
                            AuditLog.action == "download_blocked",
                        )
                    )
                ).scalars().all()
            blocked_after = (blocked[0].after_data or {}) if blocked else {}
            check_true(
                "拒绝下载留下了 download_blocked 审计（含可告警的原因）",
                len(blocked) >= 1
                and "存档" in str(blocked_after.get("reason") or ""),
                f"{[row.after_data for row in blocked]}",
            )
            rebuilt, rebuilt_headers = _download_raw(
                doc_id, token, query="allow_rebuild=true"
            )
            check(
                "明确要求时才给副本，且标明是重建的",
                rebuilt_headers.get("X-Doc-Content-Source"),
                "rebuilt-from-snapshot",
            )
            check_true("重建副本与原件不是同一份字节", rebuilt != first_bytes)
            rebuilt_texts = _cells(load_workbook(BytesIO(rebuilt)).active)
            check_true(
                "重建副本纸面上写明「由历史快照重建」",
                any("由历史快照重建" in text for text in rebuilt_texts),
                f"{rebuilt_texts}",
            )
            # `call()` 会按 JSON 解析响应体，这里下的是**二进制**，必须走 _download_raw
            try:
                state_bytes_when_missing, _h = _download_raw(
                    doc_id, token, query="mode=state"
                )
                check_true(
                    "状态副本不依赖存档，仍可下载",
                    state_bytes_when_missing[:2] == b"PK",
                    f"{state_bytes_when_missing[:2]!r}",
                )
            except AssertionError as exc:  # pragma: no cover - 只在回归失败时走到
                check_true("状态副本不依赖存档，仍可下载", False, str(exc))
        finally:
            # 复原盘上的原件：后面的清理与其它套件都假定数据是完整的
            archived_path.write_bytes(backup)
        restored, _headers = _download_raw(doc_id, token)
        check_true("复原后又能读到原件", restored == first_bytes)

        print("=== 11. §8.9 生成幂等：同一把请求键只出一份，换新键才出新版 ===")
        # 显式指定"默认模板"：第 5 节故意建了一版引用 {{order.payment_terms}} 的
        # quote_sheet 模板，它比默认模板新，不指定就会**恰好**选中它而必然是 422。
        templates = request("GET", "/biz-doc-templates?doc_type=quote_sheet", token=token)["data"]
        default_template_id = next(
            row["id"] for row in templates if row["name"].startswith("默认模板")
        )
        key = f"{PREFIX}-K1"
        base_body = {"quote_version_id": version_id, "template_id": default_template_id}
        body = {**base_body, "request_key": key}
        before_count = len(
            request("GET", f"/biz-docs?doc_type=quote_sheet&quote_id={quote_id}", token=token)["data"]
        )
        first = request("POST", "/biz-docs/quote", token=token, body=body)["data"]
        # 响应丢了、用户又点一次：带的是同一把键 → 返回**原来那一份**，不多出文件
        retry_status, retry = call("POST", "/biz-docs/quote", token=token, body=body)
        check("重发仍然 200（不是 4xx/5xx）", retry_status, 200)
        check("重发返回的是同一份单据", retry["data"]["id"], first["id"])
        check("重发没有把版本号推高", retry["data"]["version"], first["version"])
        check_true("响应明确说明这是刚才那次的重发", "未重复生成" in str(retry["message"]), f"{retry['message']}")
        after_count = len(
            request("GET", f"/biz-docs?doc_type=quote_sheet&quote_id={quote_id}", token=token)["data"]
        )
        check("台账上没有多出任何一份文件", after_count, before_count + 1)

        # 头里带键也认（脚本/移动端重试更习惯放头）
        header_status, header_retry = call(
            "POST",
            "/biz-docs/quote",
            token=token,
            body=base_body,
            headers={"X-Request-Key": key},
        )
        check("X-Request-Key 头同样被认（200）", header_status, 200)
        check("头里的同一把键也回放原文件", header_retry["data"]["id"], first["id"])

        # 明确要新版：换一把新键，**内容一模一样也照出**
        newer = request(
            "POST",
            "/biz-docs/quote",
            token=token,
            body={**base_body, "request_key": f"{PREFIX}-K2"},
        )["data"]
        check_true("换新键就是新的一版", newer["version"] > first["version"], f"{newer['version']}")
        check("新版指向前一版（链没断）", newer["parent_id"], first["id"])

        # 同一把键换了内容（这里多带了一个自定义字段）：拒绝，而不是静默按新内容再生成一份
        conflict_status, conflict = call(
            "POST",
            "/biz-docs/quote",
            token=token,
            body={
                **base_body,
                "request_key": key,
                "extra_fields": {"note": "changed-content"},
            },
        )
        check("同键换内容被拒（409）", conflict_status, 409)
        check_true("拒绝时说明了原因", "内容" in str(conflict), f"{conflict}")

    finally:
        await cleanup()

    print()
    if FAILURES:
        print(f"FAILED {len(FAILURES)} 项：" + "、".join(FAILURES))
        return 1
    print("报价冻结 + 模板变量阻断回归 全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
