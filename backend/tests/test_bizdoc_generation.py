"""第八批 §8.7 / §8.8 的落库与生成闸门回归（实际函数 + 内存 SQLite）。

为什么用真库而不是 mock：这两条的缺陷都长在"读哪一行、写哪一列"上——
- §8.7：`build_quote_doc` 以前读**当前**报价主单/客户，`unit` 写死空串；
- §8.8：`_persist` 以前不看未解析变量，直接把 status 设成 active。

把 session 换成 mock 等于自己把行为重写一遍，测不出这些错。
本文件按交接文档 §6 的口径搭一个**内存 SQLite**：真模型、真 SQL、真落库。

`numbering.generate_for` 用 `ON CONFLICT` 是 PostgreSQL 方言（迁移里就是这么建的），
SQLite 跑不了，因此这里只对取号做替身；其余一律走真实现。
"""

import asyncio
from datetime import UTC, date, datetime
from decimal import Decimal
from unittest.mock import patch

from sqlalchemy import create_engine
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session

from app.modules.bizdoc import service as bizdoc
from app.modules.bizdoc.model import BizDocTemplate
from app.modules.customer.model import Contact, Customer
from app.modules.product.model import Product, Sku
from app.modules.quote.model import Quote, QuoteCharge, QuoteItem, QuoteVersion
from app.modules.user.model import User

_BIGINT_REGISTERED = False
_JSONB_REGISTERED = False


def _sqlite_compat() -> None:
    """让 PostgreSQL 专用的类型/方言在 SQLite 上能建表。

    只影响测试进程里的 DDL 渲染，生产 DDL 仍由 Alembic 出。
    """
    global _BIGINT_REGISTERED, _JSONB_REGISTERED
    from sqlalchemy import BigInteger
    from sqlalchemy.dialects.postgresql import JSONB

    if not _BIGINT_REGISTERED:
        # SQLite 只对 `INTEGER PRIMARY KEY` 自增；本仓库主键统一 BigInteger
        @compiles(BigInteger, "sqlite")
        def _bigint_as_integer(type_, compiler, **kw):  # noqa: ARG001
            return "INTEGER"

        _BIGINT_REGISTERED = True
    if not _JSONB_REGISTERED:
        # JSONType 在 PostgreSQL 上是 JSONB；SQLite 上没有这个类型名
        @compiles(JSONB, "sqlite")
        def _jsonb_as_json(type_, compiler, **kw):  # noqa: ARG001
            return "JSON"

        _JSONB_REGISTERED = True


def _tables():
    from app.modules.bizdoc.model import BizDoc
    from app.modules.file.model import FileRecord
    from app.modules.settings.model import NumberSequence, NumberingRule, SystemSetting

    return [
        model.__table__
        for model in (
            User,
            Customer,
            Contact,
            Product,
            Sku,
            Quote,
            QuoteVersion,
            QuoteItem,
            QuoteCharge,
            BizDocTemplate,
            BizDoc,
            SystemSetting,
            NumberingRule,
            NumberSequence,
            # §8.10 之后"生成"会把实际产出的字节登记成文件行，
            # 所以要建 files 表，否则第一次生成就撞 "no such table: files"
            FileRecord,
        )
    ]


class _AsyncSession:
    """把同步 SQLite 会话包成异步会话（被测代码只 await 这几个方法）。"""

    def __init__(self, session):
        self._session = session

    async def execute(self, stmt, *args, **kwargs):
        return self._session.execute(stmt, *args, **kwargs)

    async def get(self, entity, ident, *args, **kwargs):
        return self._session.get(entity, ident, *args, **kwargs)

    async def flush(self):
        self._session.flush()

    async def commit(self):
        self._session.commit()

    async def rollback(self):
        self._session.rollback()

    def add(self, obj):
        self._session.add(obj)


def _user(user_id=901):
    from app.core.deps import CurrentUser

    row = User(
        id=user_id,
        name="冻结回归销售",
        username=f"freeze{user_id}",
        password_hash="x",
        status="active",
        created_at=datetime.now(UTC),
        updated_at=datetime.now(UTC),
    )
    return row, CurrentUser(row, permissions={"quote:view", "quote:manage"}, roles=[], data_scope="self")


def _fixture(session):
    """一条报价版本：USD、按套计价、带运费与优惠、小数数量。"""
    user, current_user = _user()
    session.add(user)
    session.flush()
    customer = Customer(name="冻结客户-原名", owner_id=user.id, status="active")
    session.add(customer)
    session.flush()
    contact = Contact(customer_id=customer.id, name="联系人-原名")
    session.add(contact)
    session.flush()
    product = Product(name="冷冻产品")
    session.add(product)
    session.flush()
    sku = Sku(
        product_id=product.id,
        sku_code="FREEZE-1",
        name="冷冻法兰",
        unit="套",
        specification="DN50",
    )
    session.add(sku)
    session.flush()
    quote = Quote(
        quote_no="Q-FREEZE-1",
        customer_id=customer.id,
        contact_id=contact.id,
        owner_id=user.id,
        status="approved",
        created_at=datetime.now(UTC),
        updated_at=datetime.now(UTC),
        valid_until=date(2026, 12, 31),
    )
    session.add(quote)
    session.flush()
    version = QuoteVersion(
        quote_id=quote.id,
        version_no=1,
        currency="USD",
        payment_terms="30% 预付",
        delivery_terms="FOB 宁波",
        trade_terms="FOB",
        subtotal_amount=Decimal("420"),
        charge_amount=Decimal("80"),
        discount_amount=Decimal("-20"),
        total_amount=Decimal("480"),
        created_at=datetime.now(UTC),
    )
    session.add(version)
    session.flush()
    quote.current_version_id = version.id
    session.add(
        QuoteItem(
            quote_version_id=version.id,
            sku_id=sku.id,
            sku_code_snapshot=sku.sku_code,
            sku_name_snapshot=sku.name,
            spec_snapshot=sku.specification,
            quantity=Decimal("3.500"),
            quoted_price=Decimal("120"),
            unit_snapshot="套",
        )
    )
    session.add(
        QuoteCharge(
            quote_version_id=version.id,
            charge_type="logistics",
            description="运费",
            amount=Decimal("80"),
            is_discount=False,
            sort_no=1,
        )
    )
    session.add(
        QuoteCharge(
            quote_version_id=version.id,
            charge_type="discount",
            description="整单优惠",
            amount=Decimal("-20"),
            is_discount=True,
            sort_no=2,
        )
    )
    session.flush()
    return current_user, customer, contact, sku, quote, version


def _numbered():
    """取号替身：真实现走 `insert ... on conflict`（PostgreSQL 方言），SQLite 跑不了。"""
    counter = {"n": 0}

    async def _generate_for(session, code, *, model, column, now=None):
        counter["n"] += 1
        return f"BJ{counter['n']:04d}"

    return _generate_for


def _run(callback):
    """跑一个异步用例：新建内存库 + 种子夹具 + 取号替身。"""
    from app.modules.settings import numbering

    _sqlite_compat()
    engine = create_engine("sqlite+pysqlite:///:memory:")
    from app.core.base import Base

    Base.metadata.create_all(engine, tables=_tables())
    with Session(engine) as raw:
        session = _AsyncSession(raw)
        current_user, customer, contact, sku, quote, version = _fixture(raw)
        raw.commit()
        with patch.object(numbering, "generate_for", _numbered()):
            return asyncio.run(
                callback(
                    session,
                    current_user=current_user,
                    customer=customer,
                    contact=contact,
                    sku=sku,
                    quote=quote,
                    version=version,
                )
            )


def test_quote_doc_freezes_version_terms_currency_and_unit():
    """§8.7：生成时的快照要带币种/单位/条款/抬头，且金额仍旧取版本。"""

    async def _case(session, *, current_user, customer, quote, version, **_):
        doc = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        await session.commit()
        snapshot = doc.input_snapshot
        frozen = snapshot["frozen"]
        assert doc.status == "active"
        assert frozen["terms"]["currency"] == "USD"
        assert frozen["terms"]["payment_terms"] == "30% 预付"
        assert frozen["terms"]["delivery_terms"] == "FOB 宁波"
        assert frozen["terms"]["valid_until"] == "2026-12-31"
        assert frozen["header"]["customer_name"] == "冻结客户-原名"
        assert frozen["header"]["contact_name"] == "联系人-原名"
        assert snapshot["items"][0]["unit"] == "套"
        assert snapshot["items"][0]["quantity"] in ("3.5", "3.500")
        # 金额仍取版本快照（已修的金额口径不许倒退）
        assert float(snapshot["total_amount"]) == 480.0
        assert [c["label"] for c in snapshot["charges"]] == ["运费", "整单优惠"]

    _run(_case)


def test_old_version_terms_unaffected_by_current_data_changes():
    """§8.7 验收：旧版本再生成后与该版条款一致；改当前 SKU/价格/客户不改已出的旧版。

    冻结分两层，别混：
    - **版本级**（币种/单位/条款/金额）：来自 `QuoteVersion`/`QuoteItem` 行上的快照，
      改当前 SKU 与价格都不影响任何一份——这一层正是 §8.7 的主诉；
    - **文档级**（客户/联系人抬头）：生成那一刻抄进快照。历史报价没有"当时的抬头"
      可追溯，所以只能按生成时点记；一旦落进快照，之后改客户资料，**已出的那一份**
      也不会变（这是"对外承诺过的文件不能被系统悄悄改掉"的底线）。
    """

    async def _case(session, *, current_user, customer, contact, sku, quote, version, **_):
        doc_v1 = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        await session.commit()
        frozen_v1 = doc_v1.input_snapshot["frozen"]
        assert frozen_v1["header"]["customer_name"] == "冻结客户-原名"

        # 之后改：客户改名、联系人改名、SKU 的名称/规格/单位都改掉。
        # **不动报价版本行**——版本行上的 quoted_price / unit_snapshot 是那一版报出去的
        # 事实，改它等于篡改那一版，不在本轮口径里（本轮挡的是"拿当前 SKU 资料补历史"）。
        customer.name = "冻结客户-已改名"
        contact.name = "联系人-已改名"
        sku.unit = "箱"
        sku.name = "冷冻法兰-改名"
        sku.specification = "DN80"
        await session.commit()

        # 旧版本再生成一份：版本级条款/单位/名称**必须与该版一致**，不受任何编辑影响
        doc_v2 = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        await session.commit()
        assert doc_v2.version == 2
        assert doc_v2.parent_id == doc_v1.id
        frozen_v2 = doc_v2.input_snapshot["frozen"]
        for field in ("currency", "payment_terms", "delivery_terms", "trade_terms", "valid_until"):
            assert frozen_v2["terms"][field] == frozen_v1["terms"][field], (
                f"旧版本再生成后条款「{field}」必须与该版一致"
            )
        assert frozen_v2["items"][0]["unit"] == "套", "单位取版本快照，不是当前 SKU 的箱"
        assert frozen_v2["items"][0]["name"] == "冷冻法兰", "名称取版本快照，不是当前 SKU 新名"
        assert frozen_v2["items"][0]["spec"] == "DN50", "规格取版本快照，不是当前 SKU 新规格"
        assert frozen_v2["items"][0]["unit_price"] == frozen_v1["items"][0]["unit_price"]
        assert frozen_v2["items"][0]["amount"] == frozen_v1["items"][0]["amount"]

        # 已出的 V1：快照与校验值一个字没动（改客户/SKU/价格都不影响它）
        reloaded = await bizdoc.get_doc_or_404(session, doc_v1.id)
        assert reloaded.input_snapshot["frozen"] == frozen_v1
        assert reloaded.input_snapshot["customer_name"] == "冻结客户-原名"
        assert reloaded.content_sha256 == doc_v1.content_sha256

        # 再开一个**新版本**：新版本出图可以带新抬头（新文件、新生成时点），
        # 但它同样不改动旧版本已经出过的任何一份
        await session.commit()
        new_version = QuoteVersion(
            quote_id=version.quote_id,
            version_no=2,
            currency="USD",
            payment_terms="全额预付",
            delivery_terms="CIF 上海",
            trade_terms="CIF",
            subtotal_amount=Decimal("100"),
            charge_amount=Decimal("0"),
            discount_amount=Decimal("0"),
            total_amount=Decimal("100"),
            created_at=datetime.now(UTC),
        )
        session.add(new_version)
        await session.flush()
        session.add(
            QuoteItem(
                quote_version_id=new_version.id,
                sku_id=sku.id,
                sku_name_snapshot=sku.name,
                quantity=Decimal("1"),
                quoted_price=Decimal("100"),
                unit_snapshot=sku.unit,  # 新版本行落的是当时（已改）的单位：箱
            )
        )
        await session.commit()
        doc_v3 = await bizdoc.generate_quote_doc(
            session, quote_version_id=new_version.id, user=current_user
        )
        await session.commit()
        assert doc_v3.input_snapshot["frozen"]["terms"]["payment_terms"] == "全额预付"
        assert doc_v3.input_snapshot["frozen"]["header"]["customer_name"] == "冻结客户-已改名"
        # 旧版仍是"原名 + 30% 预付"，不会被新版本或改名带跑
        still_v1 = await bizdoc.get_doc_or_404(session, doc_v1.id)
        assert still_v1.input_snapshot == doc_v1.input_snapshot
        assert still_v1.input_snapshot["frozen"]["terms"]["payment_terms"] == "30% 预付"
        assert still_v1.input_snapshot["customer_name"] == "冻结客户-原名"

    _run(_case)


def test_missing_unit_snapshot_is_marked_pending_not_current_sku_unit():
    """§8.7：版本行没有单位快照时写"待核实"，**不拿当前 SKU 单位静默填成历史事实**。"""

    async def _case(session, *, current_user, sku, version, **_):
        from sqlalchemy import select

        row = (
            await session.execute(
                select(QuoteItem).where(QuoteItem.quote_version_id == version.id)
            )
        ).scalars().one()
        row.unit_snapshot = None  # 模拟历史版本行
        sku.unit = "箱"  # 当前 SKU 单位已经改了
        await session.commit()

        doc = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        await session.commit()
        assert doc.input_snapshot["items"][0]["unit"] == bizdoc.UNIT_PENDING
        assert doc.input_snapshot["items"][0]["unit"] != "箱"
        assert any(
            gap["field"].endswith(".unit") for gap in doc.input_snapshot["frozen"]["gaps"]
        )
        # 渲染出去也必须是"待核实"，不能出现当前单位
        from app.modules.bizdoc.xlsx import render_quote_xlsx

        data = await bizdoc.doc_pdf_data(session, doc)
        workbook = _active(render_quote_xlsx(data))
        texts = _texts(workbook)
        assert bizdoc.UNIT_PENDING in texts
        assert "箱" not in texts

    _run(_case)


def test_unresolved_template_variable_blocks_formal_document():
    """§8.8：来源不支持的变量（订单草稿引用正式订单字段）必须阻止作为有效对外文件。"""

    async def _case(session, *, current_user, quote, version, **_):
        template = BizDocTemplate(
            doc_type="quote_sheet",
            name="引用不存在的来源",
            version=99,
            body="付款：{{order.payment_terms}}；客户：{{customer.name}}",
            enabled=True,
            created_at=datetime.now(UTC),
        )
        session.add(template)
        await session.flush()
        from app.core.errors import AppError

        try:
            await bizdoc.generate_quote_doc(
                session,
                quote_version_id=version.id,
                user=current_user,
                template_id=template.id,
            )
            raise AssertionError("未解析的模板变量没有被拦下")
        except AppError as exc:
            assert exc.http_status == 422
            assert "order.payment_terms" in str(exc.message)
            assert "来源不可用" in str(exc.message)
        await session.rollback()

    _run(_case)


def test_valid_extra_variable_is_not_blocked():
    """§8.8：合法 extra 字段（有值）必须照常出正式文件——不能把所有变量都拦死。"""

    async def _case(session, *, current_user, version, **_):
        template = BizDocTemplate(
            doc_type="quote_sheet",
            name="带自定义字段",
            version=98,
            body="客户：{{customer.name}}；项目名：{{extra.project}}；生成日：{{today}}",
            enabled=True,
            created_at=datetime.now(UTC),
        )
        session.add(template)
        await session.flush()
        doc = await bizdoc.generate_quote_doc(
            session,
            quote_version_id=version.id,
            user=current_user,
            template_id=template.id,
            extra_fields={"project": "2026 秋季项目"},
        )
        await session.commit()
        assert doc.status == "active"
        assert "2026 秋季项目" in doc.input_snapshot["body"]
        assert "冻结客户-原名" in doc.input_snapshot["body"]
        assert "{{" not in doc.input_snapshot["body"]

    _run(_case)


def test_optional_blank_field_is_allowed_but_draft_marks_missing():
    """§8.8：可选空值与未解析变量必须分开——允许空白的字段不该被当成缺值拦下。"""

    async def _case(session, *, current_user, version, **_):
        # 空的 extra 是"明确提供、值为空"：可选字段允许留白，不该拦
        blank_template = BizDocTemplate(
            doc_type="quote_sheet",
            name="可选字段留白",
            version=97,
            body="备注：{{extra.note}}",
            enabled=True,
            created_at=datetime.now(UTC),
        )
        session.add(blank_template)
        await session.flush()
        doc = await bizdoc.generate_quote_doc(
            session,
            quote_version_id=version.id,
            user=current_user,
            template_id=blank_template.id,
            extra_fields={"note": ""},
        )
        await session.commit()
        assert doc.status == "active", "明确提供、值为空的可选字段不得被拦"
        assert doc.input_snapshot["body"] == "备注："

        # 声明了 extra.xxx 却**没提供**：这是未解析变量，不是"可选空值"
        missing_template = BizDocTemplate(
            doc_type="quote_sheet",
            name="未提供的自定义字段",
            version=96,
            body="备注：{{extra.not_provided}}",
            enabled=True,
            created_at=datetime.now(UTC),
        )
        session.add(missing_template)
        await session.flush()
        from app.core.errors import AppError

        try:
            await bizdoc.generate_quote_doc(
                session,
                quote_version_id=version.id,
                user=current_user,
                template_id=missing_template.id,
            )
            raise AssertionError("未提供的 extra 变量没有被拦下")
        except AppError as exc:
            assert exc.http_status == 422
            assert "extra.not_provided" in str(exc.message)

    _run(_case)


def test_draft_preview_is_marked_and_kept_out_of_formal_ledger():
    """§8.8：明确允许时出草稿——状态是草稿、正文不带模板语法、且不可作废。"""

    async def _case(session, *, current_user, quote, version, **_):
        from app.core.errors import AppError

        template = BizDocTemplate(
            doc_type="order_sheet",
            name="订单草稿模板",
            version=95,
            body="付款：{{order.payment_terms}}",
            enabled=True,
            created_at=datetime.now(UTC),
        )
        session.add(template)
        await session.flush()
        # 报价单入口不允许出草稿：报价是对客承诺，缺条款不能出门
        try:
            await bizdoc.generate_quote_doc(
                session,
                quote_version_id=version.id,
                user=current_user,
                template_id=template.id,
            )
            raise AssertionError("报价单入口不该允许草稿")
        except AppError:
            pass
        await session.rollback()

        built = {
            "order_draft_id": 77,
            "customer_id": 1,
            "customer": None,
            "owner_id": None,
            "source": None,
            "items": [],
            "diffs": [],
            "title_suffix": "#77",
            "sections": [],
        }
        doc = await bizdoc._persist(
            session,
            built=built,
            doc_type="order_sheet",
            template=template,
            user_id=current_user.id,
            extra_fields=None,
            source_ref=None,
            allow_draft=True,
        )
        await session.commit()
        assert doc.status == "draft"
        assert doc.input_snapshot["draft"] is True
        issues = doc.input_snapshot["token_issues"]
        assert [issue["token"] for issue in issues] == ["order.payment_terms"]
        assert issues[0]["reason"] == "source_missing"
        # 正文里绝不能印出模板语法
        assert "{{" not in doc.input_snapshot["body"]
        assert "来源不可用" in doc.input_snapshot["body"]
        # 草稿不进正式台账：序列化里带标记，且不可作废
        payload = bizdoc.serialize_doc(doc)
        assert payload["is_draft"] is True
        assert payload["status_label"] == "草稿（非正式对外文件）"
        assert payload["token_issues"]
        try:
            await bizdoc.void_doc(session, doc, reason="不该能作废草稿")
            raise AssertionError("草稿不该允许作废")
        except AppError as exc:
            assert exc.http_status == 422
            assert "草稿" in str(exc.message)

    _run(_case)


def test_template_analysis_rejects_typo_and_lists_available_variables():
    """§8.8：拼错变量在保存/预览时就被指出来，并给出可用变量清单。"""
    from app.modules.bizdoc import tokens as token_rules

    analysis = token_rules.analyze_template("客户：{{customer.nmae}}，订单：{{order.order_no}}")
    assert [item["token"] for item in analysis["unsupported"]] == ["customer.nmae"]
    assert "customer.name" in analysis["unsupported"][0]["message"]
    assert analysis["variables"] == ["customer.nmae", "order.order_no"]
    assert any("customer.name" in item for item in analysis["available"])

    preview = bizdoc.preview_template_body("{{customer.nmae}}/{{extra.x}}")
    assert preview["status"] == "has_unresolved"
    assert {item["token"] for item in preview["unresolved"]} == {"customer.nmae", "extra.x"}
    assert "{{" not in preview["body"]


def test_resolve_tokens_never_leaves_template_syntax():
    """§8.8 兜底：正式文件不得印出 `{{...}}`，三种未解析原因要分得开。"""
    from app.modules.bizdoc import tokens as token_rules

    filled, unresolved = token_rules.resolve_tokens(
        "{{customer.name}}|{{order.order_no}}|{{customer.nmae}}|{{extra.todo}}",
        {"customer": {"name": "客户 A"}},
        {},
    )
    assert "{{" not in filled and "}}" not in filled
    assert "客户 A" in filled
    reasons = {item["token"]: item["reason"] for item in unresolved}
    assert reasons == {
        "order.order_no": token_rules.REASON_SOURCE_MISSING,
        "customer.nmae": token_rules.REASON_UNSUPPORTED,
        "extra.todo": token_rules.REASON_VALUE_MISSING,
    }
    assert token_rules.find_unresolved_syntax(filled) == []


def _active(xlsx_bytes):
    from io import BytesIO

    from openpyxl import load_workbook

    return load_workbook(BytesIO(xlsx_bytes)).active


def _texts(sheet):
    return [
        str(sheet.cell(row=row, column=column).value)
        for row in range(1, sheet.max_row + 1)
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(row=row, column=column).value is not None
    ]
