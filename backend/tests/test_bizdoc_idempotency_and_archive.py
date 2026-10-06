"""第八批 §8.9（生成幂等 / 版本并发）与 §8.10（原件存档）的回归。

为什么用真函数 + 内存 SQLite，而不是 mock：
这两条缺陷都长在"读哪一行、写哪一列、给不给字节"上——
- §8.9：`_next_doc_version` 是"读最新 +1"，没有排他；生成入口没有请求键；
- §8.10：下载是"当场重新渲染"，库里只有对**渲染输入**算的内容哈希，
  没有实际产出的字节。

把 session 换成 mock 等于把行为重写一遍，测不出这些。并发本身（真锁语义）
SQLite 验不了，另见 `scripts/check_biz_doc_concurrency.py`（必须连真 PostgreSQL）。

`numbering.generate_for` 用 `ON CONFLICT`（PostgreSQL 方言），SQLite 跑不了，
所以只对取号做替身；其余一律走真实现，包括**真的把 xlsx/pdf 写到 pytest 的
tmp 目录**（生成即存档，§8.10 的核心行为就在字节上）。
"""

import asyncio
import hashlib
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pytest
from sqlalchemy import create_engine, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.modules.bizdoc import service as bizdoc
from app.modules.bizdoc.model import (
    ARCHIVE_ARCHIVED,
    ARCHIVE_LEGACY,
    ARCHIVE_MISSING,
    BizDoc,
    BizDocTemplate,
)
from app.modules.customer.model import Contact, Customer
from app.modules.product.model import Product, Sku
from app.modules.quote.model import Quote, QuoteCharge, QuoteItem, QuoteVersion
from app.modules.user.model import User

_BIGINT_REGISTERED = False
_JSONB_REGISTERED = False


def _sqlite_compat() -> None:
    """让 PostgreSQL 专用类型/方言在 SQLite 上能建表（只影响测试进程的 DDL）。"""
    global _BIGINT_REGISTERED, _JSONB_REGISTERED
    from sqlalchemy import BigInteger
    from sqlalchemy.dialects.postgresql import JSONB
    from sqlalchemy.ext.compiler import compiles

    if not _BIGINT_REGISTERED:

        @compiles(BigInteger, "sqlite")
        def _bigint_as_integer(type_, compiler, **kw):  # noqa: ARG001
            return "INTEGER"

        _BIGINT_REGISTERED = True
    if not _JSONB_REGISTERED:

        @compiles(JSONB, "sqlite")
        def _jsonb_as_json(type_, compiler, **kw):  # noqa: ARG001
            return "JSON"

        _JSONB_REGISTERED = True


def _tables(*, with_files: bool = True):
    from app.core.idempotency import RequestKey
    from app.modules.file.model import FileRecord
    from app.modules.settings.model import NumberSequence, NumberingRule, SystemSetting

    models = [
        User,
        Customer,
        Contact,
        Product,
        Sku,
        Quote,
        QuoteVersion,
        QuoteItem,
        QuoteCharge,
        BizDocTemplate,
        BizDoc,
        SystemSetting,
        NumberingRule,
        NumberSequence,
        RequestKey,
    ]
    if with_files:
        models.append(FileRecord)
    return [model.__table__ for model in models]


class _AsyncSession:
    """把同步 SQLite 会话包成异步会话（被测代码只 await 这几个方法）。

    `bind` 要暴露出来：`_lock_sequence` 靠方言名判断"有没有咨询锁"，
    替身会话若没有这个属性，那段分支就永远测不到（SQLite 上应当跳过加锁）。
    """

    def __init__(self, session):
        self._session = session
        self.bind = session.get_bind()

    async def execute(self, stmt, *args, **kwargs):
        return self._session.execute(stmt, *args, **kwargs)

    async def get(self, entity, ident, *args, **kwargs):
        return self._session.get(entity, ident, *args, **kwargs)

    async def scalar(self, stmt, *args, **kwargs):
        return self._session.scalar(stmt, *args, **kwargs)

    async def scalars(self, stmt, *args, **kwargs):
        return self._session.scalars(stmt, *args, **kwargs)

    async def flush(self):
        self._session.flush()

    async def commit(self):
        self._session.commit()

    async def rollback(self):
        self._session.rollback()

    async def refresh(self, obj):
        self._session.refresh(obj)

    def add(self, obj):
        self._session.add(obj)

    async def delete(self, obj):
        self._session.delete(obj)


@pytest.fixture(autouse=True)
def isolated_file_root(tmp_path, monkeypatch):
    """§8.10 起"生成"会真的落盘：把存储根指到 pytest 的 tmp。

    不隔离的话每次跑用例都往 `backend/data/files/` 堆文件——那是运行时数据目录，
    备份额度脚本正是拿它跟库里对账的，孤儿文件会被当成异常。
    """
    from app.core.config import settings

    root = tmp_path / "files"
    root.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(settings, "file_root", str(root))
    return root


def _fixture(session):
    """一条报价版本：USD、按套计价、带运费与优惠、小数数量。"""
    user = User(
        id=901,
        name="单据回归销售",
        username="bizdoc901",
        password_hash="x",
        status="active",
        created_at=datetime.now(UTC),
        updated_at=datetime.now(UTC),
    )
    session.add(user)
    session.flush()
    customer = Customer(name="单据回归客户", owner_id=user.id, status="active")
    session.add(customer)
    session.flush()
    contact = Contact(customer_id=customer.id, name="单据回归联系人")
    session.add(contact)
    session.flush()
    product = Product(name="单据回归产品")
    session.add(product)
    session.flush()
    sku = Sku(
        product_id=product.id,
        sku_code="DOC-1",
        name="单据回归法兰",
        unit="套",
        specification="DN50",
    )
    session.add(sku)
    session.flush()
    quote = Quote(
        quote_no="Q-DOC-1",
        customer_id=customer.id,
        contact_id=contact.id,
        owner_id=user.id,
        status="approved",
        created_at=datetime.now(UTC),
        updated_at=datetime.now(UTC),
        valid_until=date(2026, 12, 31),
    )
    session.add(quote)
    session.flush()
    version = QuoteVersion(
        quote_id=quote.id,
        version_no=1,
        currency="USD",
        payment_terms="30% 预付",
        delivery_terms="FOB 宁波",
        trade_terms="FOB",
        subtotal_amount=Decimal("420"),
        charge_amount=Decimal("80"),
        discount_amount=Decimal("-20"),
        total_amount=Decimal("480"),
        created_at=datetime.now(UTC),
    )
    session.add(version)
    session.flush()
    quote.current_version_id = version.id
    session.add(
        QuoteItem(
            quote_version_id=version.id,
            sku_id=sku.id,
            sku_code_snapshot=sku.sku_code,
            sku_name_snapshot=sku.name,
            spec_snapshot=sku.specification,
            quantity=Decimal("3.500"),
            quoted_price=Decimal("120"),
            unit_snapshot="套",
        )
    )
    session.add(
        QuoteCharge(
            quote_version_id=version.id,
            charge_type="logistics",
            description="运费",
            amount=Decimal("80"),
            is_discount=False,
            sort_no=1,
        )
    )
    session.add(
        QuoteCharge(
            quote_version_id=version.id,
            charge_type="discount",
            description="整单优惠",
            amount=Decimal("-20"),
            is_discount=True,
            sort_no=2,
        )
    )
    session.flush()
    from app.core.deps import CurrentUser

    current_user = CurrentUser(
        user, permissions={"quote:view", "quote:manage"}, roles=[], data_scope="self"
    )
    return {
        "raw": session,
        "current_user": current_user,
        "customer": customer,
        "quote": quote,
        "version": version,
        "sku": sku,
    }


def _numbered():
    """取号替身：真实现走 `insert ... on conflict`（PostgreSQL 方言）。"""
    counter = {"n": 0}

    async def _generate_for(session, code, *, model, column, now=None):
        counter["n"] += 1
        return f"BJ{counter['n']:04d}"

    return _generate_for


def run_case(case, *, with_files: bool = True, **extra):
    """跑一个异步用例：内存库 + 夹具 + 取号替身；其余走真实现。"""
    from app.modules.settings import numbering

    _sqlite_compat()
    engine = create_engine("sqlite+pysqlite:///:memory:")
    from app.core.base import Base

    Base.metadata.create_all(engine, tables=_tables(with_files=with_files))
    with Session(engine) as raw:
        session = _AsyncSession(raw)
        fixtures = _fixture(raw)
        raw.commit()
        with patch.object(numbering, "generate_for", _numbered()):
            return asyncio.run(case(session, **fixtures, **extra))


def _stored_files(root: Path) -> list[Path]:
    return [p for p in root.rglob("*") if p.is_file()]


# ---------------------------------------------------------------- §8.10 存档


def test_generation_archives_real_bytes_and_download_reads_them(isolated_file_root):
    """§8.10 验收：生成即存档；重复下载拿到**同一份字节**，不再重新渲染。

    关键差异（这条用例真正要钉住的）：
    `content_sha256` 是对**渲染输入**算的，`file_sha256` 是对**实际产出的字节**
    算的——存档之后下载根本不再调用渲染器，所以把渲染器打成"一调用就炸"，
    下载照样出得来；只有这样才能证明"读的是归档，不是又渲染了一遍"。
    """

    async def _case(session, current_user, customer, quote, version, **_):
        doc = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        await session.commit()

        # 1) 生成即存档：文件行、字节哈希、渲染器版本都落上了
        assert doc.file_id is not None
        assert doc.renderer_version == "bizdoc-xlsx/1"
        stored = _stored_files(isolated_file_root)
        assert len(stored) == 1, f"生成应当只产生一个原件，实际 {stored}"
        payload = stored[0].read_bytes()
        assert doc.file_size == len(payload)
        assert doc.file_sha256 == hashlib.sha256(payload).hexdigest()
        # 内容哈希与字节哈希是**两件事**：一个证明输入没变，一个证明字节没变
        assert doc.content_sha256 != doc.file_sha256
        assert bizdoc.archive_status_of(doc) == ARCHIVE_ARCHIVED

        # 2) 下载走归档：把渲染器换成"一调用就炸"，原样读得到
        def _explode(*_args, **_kwargs):  # pragma: no cover - 被调用即失败
            raise AssertionError("下载存档原件时不该重新渲染")

        with patch.object(bizdoc, "render_document", _explode):
            first = await bizdoc.plan_download(session, doc)
            second = await bizdoc.plan_download(session, doc)
        assert first.content == payload and second.content == payload
        assert first.source == bizdoc.SOURCE_ARCHIVED
        assert first.blocked_reason is None
        # 3) 重复下载字节一致（验收原话）
        assert hashlib.sha256(first.content).hexdigest() == hashlib.sha256(
            second.content
        ).hexdigest()

    run_case(_case)


def test_state_copy_is_always_rerendered_and_marked(isolated_file_root):
    """§8.10：作废件的"已作废"看**状态副本**，那份不是原件，必须标明。"""

    async def _case(session, current_user, version, **_):
        doc = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        await session.commit()
        await bizdoc.void_doc(session, doc, reason="回归作废")
        await session.commit()

        plan = await bizdoc.plan_download(session, doc, mode="state")
        assert plan.content is None, "状态副本一定是重出的，不能拿原件顶"
        assert plan.source == bizdoc.SOURCE_STATE_COPY
        assert plan.notice == bizdoc.STATE_COPY_NOTICE

        data = await bizdoc.doc_pdf_data(session, doc, copy_notice=plan.notice)
        assert data["copy_notice"] == bizdoc.STATE_COPY_NOTICE
        assert data["status_label"] == "已作废"
        # 纸面上真的印出来了吗？Excel 可以回读单元格；PDF 走"加了标记字节就不同"
        rendered = bizdoc.render_document(doc.doc_type, data)
        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(rendered)).active
        texts = [
            str(sheet.cell(row=row, column=column).value)
            for row in range(1, sheet.max_row + 1)
            for column in range(1, sheet.max_column + 1)
            if sheet.cell(row=row, column=column).value is not None
        ]
        assert any(bizdoc.STATE_COPY_NOTICE in text for text in texts), texts

        plain = bizdoc.render_document(
            doc.doc_type, await bizdoc.doc_pdf_data(session, doc)
        )
        assert plain != rendered, "副本标记必须真的印进文件里"

    run_case(_case)


def test_legacy_doc_without_archive_is_marked_rebuilt(isolated_file_root):
    """§8.10：历史没有存档的文件标明"由历史快照重建"，不伪称当时的原件。"""

    async def _case(session, current_user, version, **_):
        doc = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        await session.commit()
        # 模拟迁移前生成的历史行：没有 file_id / 字节哈希 / 渲染器版本
        doc.file_id = None
        doc.file_sha256 = None
        doc.file_size = None
        doc.renderer_version = None
        await session.commit()

        assert bizdoc.archive_status_of(doc) == ARCHIVE_LEGACY
        plan = await bizdoc.plan_download(session, doc)
        assert plan.content is None
        assert plan.source == bizdoc.SOURCE_REBUILT
        assert plan.notice == bizdoc.REBUILD_NOTICE
        payload = bizdoc.serialize_doc(doc)
        assert payload["archive"]["status"] == ARCHIVE_LEGACY
        # 重建件上必须写明它不是原件
        data = await bizdoc.doc_pdf_data(session, doc, copy_notice=plan.notice)
        assert data["copy_notice"] == bizdoc.REBUILD_NOTICE

    run_case(_case)


def test_missing_or_corrupt_archive_is_blocked_not_silently_substituted(
    isolated_file_root,
):
    """§8.10 验收：失踪/损坏原件要告警；**不静默用当前资料替代**。

    这一条是整批里最容易被"顺手回退成重新渲染"糊过去的地方：重新渲染永远成功，
    所以只有明确断言"被拒绝"才守得住。
    """
    from app.modules.file.model import FileRecord

    async def _case(session, current_user, version, **_):
        doc = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        await session.commit()
        record = await session.get(FileRecord, doc.file_id)

        # ① 盘上的原件被删掉：拒绝下载（不是"重新渲染一份"）
        from app.modules.file import storage

        path = storage.absolute_path(record.object_key)
        path.unlink()
        plan = await bizdoc.plan_download(session, doc)
        assert plan.content is None
        assert plan.blocked_reason, "原件不见了必须给出可告警的原因"
        assert plan.source == ARCHIVE_MISSING
        # ② 明确要求时，才给一份**标注过**的重建副本
        allowed = await bizdoc.plan_download(session, doc, allow_rebuild=True)
        assert allowed.blocked_reason is None
        assert allowed.notice == bizdoc.REBUILD_NOTICE
        assert allowed.source == bizdoc.SOURCE_REBUILT

        # ③ 文件在但字节被换过：校验值对不上，同样拒绝
        path.write_bytes(b"%PDF-1.4 replaced-original-bytes")
        plan = await bizdoc.plan_download(session, doc)
        assert plan.content is None
        assert "校验值" in plan.blocked_reason
        # ④ 文件登记记录整个不见了
        await session.delete(record)
        await session.commit()
        plan = await bizdoc.plan_download(session, doc)
        assert plan.content is None
        assert plan.blocked_reason

    run_case(_case)


def test_failed_archive_registration_leaves_no_doc_and_no_orphan_file(
    isolated_file_root,
):
    """§8.10：生成/文件登记失败要能补偿——不留"active 却没有原件"的单据。

    这里的失败是**文件登记**失败（测试库故意不建 `files` 表，等价于登记这一步
    出错）：写盘已经发生，但整条业务行必须回滚，盘上那份半成品必须被删掉。
    """
    from app.modules.file import storage  # noqa: F401  （确保模块已加载）

    async def _case(session, raw, current_user, version, **_):
        with pytest.raises(Exception):
            await bizdoc.generate_quote_doc(
                session, quote_version_id=version.id, user=current_user
            )
        raw.rollback()
        # 没有虚假的有效单据
        assert raw.scalar(select(func.count()).select_from(BizDoc)) == 0
        # 也没有谁也认领不了的孤儿原件
        assert _stored_files(isolated_file_root) == []

    run_case(_case, with_files=False)


def test_commit_failure_deletes_the_archived_object(isolated_file_root):
    """§8.10：提交失败时把这次落盘的原件删掉（否则盘上留下无记录的文件）。"""

    async def _case(session, raw, current_user, version, **_):
        doc = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        assert len(_stored_files(isolated_file_root)) == 1

        async def _boom():
            raise RuntimeError("提交失败")

        with patch.object(session, "commit", _boom):
            with pytest.raises(RuntimeError):
                await bizdoc.commit_doc_generation(session, doc)
        assert _stored_files(isolated_file_root) == []

    run_case(_case)


def test_renderer_version_comes_from_the_renderer_module():
    """§8.10：渲染器版本只有一处定义——升级渲染器时版本号不会漏改。"""
    from app.modules.bizdoc import pdf as pdf_renderer
    from app.modules.bizdoc import xlsx as xlsx_renderer

    assert bizdoc.renderer_spec("quote_sheet")[3] == xlsx_renderer.RENDERER_VERSION
    assert bizdoc.renderer_spec("order_sheet")[3] == pdf_renderer.RENDERER_VERSION
    assert bizdoc.renderer_spec("sample_request")[0] == ".pdf"


# ------------------------------------------------- §8.9 来源链 / 唯一约束


def test_source_key_separates_order_draft_from_order_and_quote_from_others():
    """§8.9：来源类型与 NULL 的语义——**不同来源不许共用历史链**。"""
    assert bizdoc.source_key_for(doc_type="sample_request", sample_request_id=3) == "sample_request:3"
    assert bizdoc.source_key_for(doc_type="order_sheet", order_id=4) == "order:4"
    # 草稿与它转出来的正式订单是两条链（草稿 V3 不该把正式订单顶到 V4）
    assert bizdoc.source_key_for(
        doc_type="order_sheet", order_id=4, order_draft_id=5
    ) == "order_draft:5"
    assert bizdoc.source_key_for(doc_type="quote_sheet", quote_id=6) == "quote:6"
    # 取不到来源 → 不参与同源唯一约束（历史无法归属的行）
    assert bizdoc.source_key_for(doc_type="order_sheet") is None
    # 同一类型的两个不同来源必然是不同的键
    assert bizdoc.source_key_for(
        doc_type="sample_request", sample_request_id=1
    ) != bizdoc.source_key_for(doc_type="sample_request", sample_request_id=2)


def test_same_source_versions_are_sequential_and_isolated(isolated_file_root):
    """§8.9 验收：两次明确新版连续且 parent 链正确；不同来源各有各的链。"""

    async def _case(session, raw, current_user, version, **_):
        doc1 = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        await session.commit()
        doc2 = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        await session.commit()
        assert (doc1.version, doc1.parent_id) == (1, None)
        assert (doc2.version, doc2.parent_id) == (2, doc1.id)
        assert doc1.source_key == doc2.source_key == f"quote:{version.quote_id}"

        # 另一个来源（另一张报价单）必须从 V1 起，不能蹭上一条链
        other_quote = Quote(
            quote_no="Q-DOC-2",
            customer_id=doc1.customer_id,
            owner_id=current_user.id,
            status="approved",
            created_at=datetime.now(UTC),
            updated_at=datetime.now(UTC),
        )
        raw.add(other_quote)
        raw.flush()
        other_version = QuoteVersion(
            quote_id=other_quote.id,
            version_no=1,
            currency="CNY",
            subtotal_amount=Decimal("10"),
            charge_amount=Decimal("0"),
            discount_amount=Decimal("0"),
            total_amount=Decimal("10"),
            created_at=datetime.now(UTC),
        )
        raw.add(other_version)
        raw.flush()
        other_doc = await bizdoc.generate_quote_doc(
            session, quote_version_id=other_version.id, user=current_user
        )
        await session.commit()
        assert other_doc.version == 1
        assert other_doc.parent_id is None
        assert other_doc.source_key != doc1.source_key

    run_case(_case)


def test_unique_index_rejects_duplicate_source_version(isolated_file_root):
    """§8.9：**数据库层**必须挡住"同源同版本"——代码写错也造不出重复链。"""

    async def _case(session, raw, current_user, version, **_):
        doc = await bizdoc.generate_quote_doc(
            session, quote_version_id=version.id, user=current_user
        )
        await session.commit()
        clone = BizDoc(
            doc_no="BJ9999",
            doc_type=doc.doc_type,
            title="重复版本",
            version=doc.version,  # 同一个来源的同一个版本
            source_key=doc.source_key,
            status="active",
            template_id=doc.template_id,
            template_version=doc.template_version,
            input_snapshot={},
            content_sha256="x" * 64,
            created_at=datetime.now(UTC),
        )
        raw.add(clone)
        with pytest.raises(IntegrityError):
            raw.flush()
        raw.rollback()

        # 反过来：来源键为空（历史无法归属）的行**不该**互相判冲突——
        # 否则迁移时一堆积压的历史行会直接卡死整个迁移
        for index in range(2):
            raw.add(
                BizDoc(
                    doc_no=f"BJ88{index}",
                    doc_type="order_sheet",
                    title="无来源历史件",
                    version=1,
                    source_key=None,
                    status="active",
                    template_id=doc.template_id,
                    template_version=1,
                    input_snapshot={},
                    content_sha256="y" * 64,
                    created_at=datetime.now(UTC),
                )
            )
        raw.flush()  # 不抛异常才算通过
        raw.rollback()

    run_case(_case)


# ---------------------------------------------------------------- §8.9 幂等


def test_same_request_key_replays_and_new_key_allows_new_version(isolated_file_root):
    """§8.9 验收：双击/响应丢失只一份；明确要新版换新键就照出（哪怕内容一样）。"""

    async def _case(session, raw, current_user, version, **_):
        calls = {"n": 0}

        def _generate():
            calls["n"] += 1
            return bizdoc.generate_quote_doc(
                session, quote_version_id=version.id, user=current_user
            )

        payload = {"quote_version_id": version.id}
        first, replayed = await bizdoc.run_idempotent_generation(
            session,
            user_id=current_user.id,
            action="bizdoc:quote_sheet",
            request_key="key-1",
            payload=payload,
            generate=_generate,
        )
        await session.commit()
        assert replayed is False and calls["n"] == 1

        # 同一把键重发（响应丢了、用户又点了一次）：返回**原来那份**，不新建
        again, replayed = await bizdoc.run_idempotent_generation(
            session,
            user_id=current_user.id,
            action="bizdoc:quote_sheet",
            request_key="key-1",
            payload=payload,
            generate=_generate,
        )
        await session.commit()
        assert replayed is True
        assert again.id == first.id
        assert calls["n"] == 1, "回放不该再调用一次生成"
        assert raw.scalar(select(func.count()).select_from(BizDoc)) == 1

        # 明确要新版：换一把新键。**内容完全一样也必须出新版**——
        # 按内容去重会把"同一批货再出一份给工厂"这种合法需求永久挡掉。
        third, replayed = await bizdoc.run_idempotent_generation(
            session,
            user_id=current_user.id,
            action="bizdoc:quote_sheet",
            request_key="key-2",
            payload=payload,
            generate=_generate,
        )
        await session.commit()
        assert replayed is False
        assert calls["n"] == 2
        assert third.version == 2 and third.parent_id == first.id
        # 内容一模一样也照出新版（`content_sha256` 含单号，所以两份的**内容哈希
        # 不同是正常的**；要证明的是"快照一字不差，但确实出了新的一版"）
        assert third.input_snapshot == first.input_snapshot
        assert third.doc_no != first.doc_no

        # 同一把键换了内容：报冲突，不静默按新内容再生成一份
        from app.core.errors import AppError

        with pytest.raises(AppError) as excinfo:
            await bizdoc.run_idempotent_generation(
                session,
                user_id=current_user.id,
                action="bizdoc:quote_sheet",
                request_key="key-2",
                payload={"quote_version_id": 999999},
                generate=_generate,
            )
        assert excinfo.value.http_status == 409
        raw.rollback()
        assert raw.scalar(select(func.count()).select_from(BizDoc)) == 2

    run_case(_case)


def test_generation_without_request_key_still_works_and_is_not_faked(
    isolated_file_root,
):
    """§8.9：不给键 = 不做幂等（不假装有保护），照常生成。"""

    async def _case(session, current_user, version, **_):
        doc, replayed = await bizdoc.run_idempotent_generation(
            session,
            user_id=current_user.id,
            action="bizdoc:quote_sheet",
            request_key=None,
            payload={},
            generate=lambda: bizdoc.generate_quote_doc(
                session, quote_version_id=version.id, user=current_user
            ),
        )
        await session.commit()
        assert replayed is False and doc.id is not None

    run_case(_case)


def test_released_placeholder_allows_retry_with_same_key(isolated_file_root):
    """§8.9：生成失败要释放占位——用户改完资料带着同一把键重试不能被判成冲突。"""

    async def _case(session, raw, current_user, version, **_):
        from app.core.errors import AppError, ErrorCode

        async def _fail():
            raise AppError(ErrorCode.PARAM_ERROR, "模板变量未解析", 422)

        with pytest.raises(AppError):
            await bizdoc.run_idempotent_generation(
                session,
                user_id=current_user.id,
                action="bizdoc:quote_sheet",
                request_key="retry-key",
                payload={"quote_version_id": version.id},
                generate=_fail,
            )
        raw.rollback()

        # 改完资料重试：同一把键、**不同内容**，必须能正常生成
        doc, replayed = await bizdoc.run_idempotent_generation(
            session,
            user_id=current_user.id,
            action="bizdoc:quote_sheet",
            request_key="retry-key",
            payload={"quote_version_id": version.id, "fixed": True},
            generate=lambda: bizdoc.generate_quote_doc(
                session, quote_version_id=version.id, user=current_user
            ),
        )
        await session.commit()
        assert replayed is False and doc.version == 1

    run_case(_case)


# ---------------------------------------------------------------- §8.9 模板


def test_template_version_conflict_is_explained_not_500(isolated_file_root):
    """§8.9：模板版本并发不能直接 500——唯一冲突要翻成一句能照着处理的 409。"""

    async def _case(session, raw, **_):
        from app.core.errors import AppError

        session.add(
            BizDocTemplate(
                doc_type="quote_sheet",
                name="已有模板",
                version=1,
                body="",
                enabled=True,
                created_at=datetime.now(UTC),
            )
        )
        await session.commit()
        # 模拟并发：另一个连接已经占掉了 V1，而这次仍以为下一个是 V1
        async def _stale_next_version(_session, _doc_type):
            return 1

        with patch.object(bizdoc, "next_template_version", _stale_next_version):
            with pytest.raises(AppError) as excinfo:
                await bizdoc.create_template_version(
                    session,
                    doc_type="quote_sheet",
                    name="并发模板",
                    body="",
                    enabled=True,
                    remark=None,
                    user_id=1,
                )
        assert excinfo.value.http_status == 409
        assert "V1" in excinfo.value.message
        # 半成品被回滚干净，**已存在的那一版不受影响**，调用方还能接着读
        assert (
            raw.scalar(
                select(func.count())
                .select_from(BizDocTemplate)
                .where(BizDocTemplate.doc_type == "quote_sheet")
            )
            == 1
        )
        assert (await bizdoc.next_template_version(session, "quote_sheet")) == 2

    run_case(_case)


def test_default_template_seeding_is_idempotent(isolated_file_root):
    """§8.9：冷启动补默认模板要幂等（并发时靠"锁后再复查一次"）。"""

    async def _case(session, raw, **_):
        await bizdoc.ensure_default_templates(session)
        await session.commit()
        first = raw.scalar(select(func.count()).select_from(BizDocTemplate))
        await bizdoc.ensure_default_templates(session)
        await session.commit()
        assert raw.scalar(select(func.count()).select_from(BizDocTemplate)) == first
        assert first == len(bizdoc.DOC_TYPE_LABEL)

    run_case(_case)
