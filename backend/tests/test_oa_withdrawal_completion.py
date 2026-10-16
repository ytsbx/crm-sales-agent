"""撤回真实外部动作的错误分支、最终快照补收及完整打样资料回归。"""
import io
from datetime import UTC, datetime
from types import SimpleNamespace as Obj
from unittest.mock import AsyncMock

import pytest
import httpx
from openpyxl import load_workbook
from starlette.requests import Request

from app.core.config import settings
from app.core.errors import AppError
from app.modules.dingtalk import router, service
from app.modules.dingtalk.client import DingTalkError, DingTalkUnknownOutcome
from app.modules.dingtalk.model import OaInstance
from app.modules.sample.application_form import render_application_form
from app.modules.sample.oa_content import submission_summary


@pytest.fixture
def withdrawal(monkeypatch):
    monkeypatch.setattr(settings, 'dingtalk_push_off', False)
    row=OaInstance(id=1,inquiry_id=7,inquiry_version=1,oa_type='sample',
                   instance_id='ding-7',status='pending',result={'payload':{'tasks':[]}})
    account=Obj(dingtalk_userid='actor')
    session=Obj(get=AsyncMock(side_effect=[row,account]),commit=AsyncMock())
    client=Obj(terminate_process_instance=AsyncMock(return_value={'result':True}))
    from app.modules.dingtalk import client as module
    monkeypatch.setattr(module,'get_client',lambda:client)
    monkeypatch.setattr(router,'ensure_oa_type_scope',AsyncMock())
    monkeypatch.setattr(router,'write_audit',AsyncMock())
    req=Request({'type':'http','headers':[],'client':('127.0.0.1',1)})
    return row,session,client,req


@pytest.mark.anyio
async def test_withdraw_marks_final_snapshot_for_collection(withdrawal):
    row,session,client,req=withdrawal
    await router.withdraw_oa_instance(1,req,Obj(id=1,name='销售'),session)
    assert row.status=='withdrawn'
    assert row.result['withdraw_snapshot_pending'] is True
    assert row.result['payload']=={'tasks':[]}
    client.terminate_process_instance.assert_awaited_once_with(
        instance_id='ding-7',operating_user_id='actor',remark='CRM 撤回钉钉审批')


@pytest.mark.anyio
@pytest.mark.parametrize('error',[
    DingTalkUnknownOutcome('timeout'),DingTalkError('permission denied',api='terminate'),
])
async def test_failed_or_unknown_withdrawal_never_claims_withdrawn(withdrawal,error):
    row,session,client,req=withdrawal
    client.terminate_process_instance.side_effect=error
    with pytest.raises(AppError):
        await router.withdraw_oa_instance(1,req,Obj(id=1,name='销售'),session)
    assert row.status=='pending'
    assert not row.result.get('crm_withdrawn')
    if isinstance(error,DingTalkUnknownOutcome):
        assert '撤回结果暂未确认' in row.error
        session.commit.assert_awaited_once()


@pytest.mark.anyio
@pytest.mark.parametrize('status',['RUNNING','TERMINATED','CANCELED','CANCELLED'])
async def test_withdraw_final_snapshot_does_not_regress_to_pending(monkeypatch,status):
    monkeypatch.setattr(settings,'dingtalk_push_off',False)
    row=OaInstance(id=1,inquiry_id=7,inquiry_version=1,oa_type='inquiry',instance_id='ding-7',
                   status='withdrawn',result={'crm_withdrawn':True,'withdraw_snapshot_pending':True})
    session=Obj(execute=AsyncMock(return_value=Obj(scalars=lambda:Obj(all=lambda:[row]))),flush=AsyncMock())
    payload={'status':status,'operationRecords':[{'type':'TERMINATE_PROCESS_INSTANCE','remark':'已撤回'}]}
    monkeypatch.setattr(service,'get_client',lambda:Obj(get_process_instance=AsyncMock(return_value=payload)))
    await service.sync_pending_instances(session)
    assert row.status=='withdrawn'
    assert row.result['withdraw_snapshot_pending'] is (status=='RUNNING')
    if status=='RUNNING':
        assert 'payload' not in row.result  # 旧流转快照不得覆盖已确认的撤回。
    else:
        assert row.result['payload']==payload
    assert row.result['crm_withdrawn']


@pytest.mark.anyio
async def test_query_collects_withdrawn_snapshots_but_leaves_completed_history_alone(monkeypatch):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from tests.conftest import SyncSessionAsAsync, _register_sqlite_bigint_as_integer
    _register_sqlite_bigint_as_integer()
    monkeypatch.setattr(settings,'dingtalk_push_off',False)
    engine=create_engine('sqlite://')
    OaInstance.__table__.create(engine)
    with Session(engine) as raw:
        for index,(state,result) in enumerate([
            ('pending',None),('withdrawn',{'withdraw_snapshot_pending':True}),
            ('withdrawn',{'withdraw_snapshot_pending':False}),('approved',{}),
        ],1):
            raw.add(OaInstance(id=index,inquiry_id=index,inquiry_version=1,oa_type='inquiry',
                instance_id=f'ding-{index}',idempotency_key=f'key-{index}',status=state,result=result,
                created_at=datetime.now(UTC)))
        raw.commit()
        read=AsyncMock(return_value={'status':'TERMINATED'})
        monkeypatch.setattr(service,'get_client',lambda:Obj(get_process_instance=read))
        await service.sync_pending_instances(SyncSessionAsAsync(raw))
        assert {call.args[0] for call in read.await_args_list}=={'ding-1','ding-2'}
    engine.dispose()


def test_sample_keeps_each_lines_material_craft_and_drawing_in_body_and_excel():
    sample=Obj(id=7,product_name='两种盒子',material_spec='单头材质')
    items=[Obj(item_name='纸盒',quantity=2,material='纸',craft='压纹',drawing_version='A'),
           Obj(item_name='塑料盒',quantity=3,material='PP',craft='热压',drawing_version='B')]
    summary=submission_summary(sample,items,note='逐项制作')
    for text in ('逐项制作','纸盒','塑料盒','数量：2','数量：3','材质：PP','工艺：压纹','图纸版本：B'):
        assert text in summary
    sample.oa_submission_summary=summary
    book=load_workbook(io.BytesIO(render_application_form(sample,items)))
    text='\n'.join(str(row[0] or '') for row in book['提交资料与逐条明细'].values)
    assert text==summary


def test_excel_summary_preserves_long_text_and_formula_like_input():
    long='详'*40000
    sample=Obj(id=7,oa_submission_summary='='+long)
    book=load_workbook(io.BytesIO(render_application_form(sample,[])))
    cells=list(book['提交资料与逐条明细']['A'])
    assert ''.join(cell.value for cell in cells)=='='+long
    assert all(cell.data_type=='s' for cell in cells)


@pytest.mark.anyio
@pytest.mark.parametrize('kind',['inquiry','sample','contract'])
async def test_view_only_user_cannot_withdraw_or_resend(kind):
    with pytest.raises(AppError):
        await router.ensure_oa_type_scope(Obj(),Obj(roles=[],has=lambda _:False),Obj(oa_type=kind))


@pytest.mark.anyio
async def test_sample_management_does_not_require_quote_management(monkeypatch):
    from app.modules.sample import service as samples
    visible=AsyncMock()
    monkeypatch.setattr(samples,'get_visible_or_404',visible)
    actor=Obj(roles=[],has=lambda code:code=='sample:manage')
    session=Obj()
    await router.ensure_oa_type_scope(session,actor,Obj(oa_type='sample',inquiry_id=7))
    visible.assert_awaited_once_with(session,actor,7,allow_external_flow=True)


@pytest.mark.anyio
@pytest.mark.parametrize('status,payload,exception',[
    (200,{'success':True,'result':True},None),
    (200,{'result':True},None),
    (200,{'success':False},DingTalkError),
    (200,{},DingTalkUnknownOutcome),
    (200,[],DingTalkUnknownOutcome),
    (503,{},DingTalkUnknownOutcome),
])
async def test_terminate_requires_explicit_acknowledgement(monkeypatch,status,payload,exception):
    from app.modules.dingtalk.client import DingTalkClient
    def reply(request):
        assert request.url.path=='/v1.0/workflow/processInstances/terminate'
        return httpx.Response(status,json=payload)
    http=httpx.AsyncClient(transport=httpx.MockTransport(reply))
    monkeypatch.setattr(httpx,'AsyncClient',lambda **_:http)
    client=DingTalkClient()
    monkeypatch.setattr(client,'_require',lambda:None)
    monkeypatch.setattr(client,'access_token',AsyncMock(return_value='test-token'))
    if exception:
        with pytest.raises(exception):
            await client.terminate_process_instance(instance_id='test-id',operating_user_id='actor',remark='test')
    else:
        assert await client.terminate_process_instance(instance_id='test-id',operating_user_id='actor',remark='test')==payload
