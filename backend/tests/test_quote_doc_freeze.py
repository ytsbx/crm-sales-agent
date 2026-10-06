"""第八批 §8.7：对客报价 Excel 必须按报价版本冻结并展示币种/单位/条款。

复现（修前）：给实际 xlsx renderer 传入 `currency="USD"`，`render_quote_xlsx`
里没有任何一处读它——回读单元格找不到币种字段；`build_quote_doc` 的明细
`unit` 固定为空串；条款只有一行"有效期至"，且取自当前报价主单。

本文件用**实际函数 + 回读**证明修后行为（交接文档 §6 的既有证据方式）：
`build_quote_doc` 的真实输出 → 真实 `render_quote_xlsx` → openpyxl 回读。

只测只读组装与渲染，不连数据库：session 用最小替身喂固定记录。
"""

import asyncio
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

from openpyxl import load_workbook

from app.modules.bizdoc.service import (
    NOT_RETAINED,
    UNIT_PENDING,
    build_quote_doc,
    doc_pdf_data,
)
from app.modules.bizdoc.xlsx import render_quote_xlsx
from app.modules.customer.model import Contact, Customer
from app.modules.quote.model import Quote, QuoteVersion
from app.modules.user.model import User

VERSION_ID = 41
QUOTE_ID = 7
CUSTOMER_ID = 3
CONTACT_ID = 9
OWNER_ID = 5
SKU_ID = 88


def _records(*, unit_snapshot: str | None = "套", currency: str = "USD"):
    quote = SimpleNamespace(
        id=QUOTE_ID,
        quote_no="Q-CHK-0001",
        customer_id=CUSTOMER_ID,
        contact_id=CONTACT_ID,
        owner_id=OWNER_ID,
        valid_until=None,
    )
    version = SimpleNamespace(
        id=VERSION_ID,
        quote_id=QUOTE_ID,
        version_no=2,
        currency=currency,
        payment_terms=None,  # 版本上没填：必须显示"未留存"，不能空着
        delivery_terms="FOB 宁波",
        trade_terms="FOB",
        exchange_rate_snapshot=None,
        subtotal_amount=Decimal("420"),
        charge_amount=Decimal("80"),
        discount_amount=Decimal("-20"),
        total_amount=Decimal("480"),
    )
    item = SimpleNamespace(
        id=1,
        quote_version_id=VERSION_ID,
        sku_id=SKU_ID,
        inquiry_id=None,
        inquiry_no_snapshot=None,
        sku_code_snapshot="SKU-88",
        sku_name_snapshot="不锈钢法兰",
        spec_snapshot="DN50",
        quantity=Decimal("3.500"),
        quoted_price=Decimal("120"),
        unit_snapshot=unit_snapshot,
        remark=None,
    )
    return quote, version, item


def _charges():
    return [
        SimpleNamespace(
            charge_type="logistics", description="运费", amount=Decimal("80"), is_discount=False
        ),
        SimpleNamespace(
            charge_type="discount", description="整单优惠", amount=Decimal("-20"), is_discount=True
        ),
    ]


def _session(quote, version, item, charge_rows, *, customer_name="虚构客户", contact_name="王经理"):
    """最小替身：get 按模型返回记录，execute 依次返回明细行与费用行。"""
    typed = {
        (Quote, QUOTE_ID): quote,
        (QuoteVersion, VERSION_ID): version,
        (Customer, CUSTOMER_ID): SimpleNamespace(id=CUSTOMER_ID, name=customer_name),
        (Contact, CONTACT_ID): SimpleNamespace(id=CONTACT_ID, name=contact_name),
        (User, OWNER_ID): SimpleNamespace(id=OWNER_ID, name="李销售"),
    }
    item_result = Mock(scalars=Mock(return_value=Mock(all=Mock(return_value=[item]))))
    charge_result = Mock(scalars=Mock(return_value=Mock(all=Mock(return_value=charge_rows))))
    return SimpleNamespace(
        get=AsyncMock(side_effect=lambda model, ident: typed[(model, ident)]),
        execute=AsyncMock(side_effect=[item_result, charge_result]),
    )


def _build(session, version_id=VERSION_ID):
    return asyncio.run(build_quote_doc(session, version_id))


def _cells(worksheet):
    """回读整张表：[(行, 列, 文本)]，供"某一格必须有币种"这类断言使用。"""
    rows = []
    for row in range(1, worksheet.max_row + 1):
        for column in range(1, worksheet.max_column + 1):
            value = worksheet.cell(row=row, column=column).value
            if value is not None:
                rows.append((row, column, str(value)))
    return rows


def _xlsx(built):
    """把组装结果喂给**实际渲染器**再回读（交接文档 §6 的证据方式）。"""
    return load_workbook(BytesIO(render_quote_xlsx(built))).active


def _texts(sheet):
    return [text for _row, _column, text in _cells(sheet)]


def _build_and_render(**kwargs):
    quote, version, item = _records(**kwargs)
    built = _build(_session(quote, version, item, _charges()))
    return built, _texts(_xlsx(built))


def test_quote_excel_shows_version_currency_and_unit():
    """修前：输入 USD 后单元格既没有币种字段也没有 USD；单位列固定空白。"""
    _built, texts = _build_and_render()

    assert "币种" in texts, f"Excel 上必须有一栏明确的币种：{texts}"
    assert "USD" in texts, f"美元单必须印出 USD（不能只有金额）：{texts}"

    # 数量与单位同列组：表头里"数量"与"单位"相邻，数据行单位取版本快照
    headers = [text for text in texts if text in ("数量", "单位")]
    assert "单位" in headers and headers[headers.index("单位") - 1] == "数量", (
        f"数量与单位必须同列组相邻：{headers}"
    )
    assert "套" in texts, f"计价单位取版本快照（套），不能空白：{texts}"
    # 小数数量如实出现（3.5 不能印成 4，也不能被抹成 3）
    assert "3.500" in texts or "3.5" in texts, f"小数数量必须如实：{texts}"


def test_quote_excel_shows_terms_totals_and_header():
    """条款、抬头、附加费/优惠、合计都要在表上，且都来自版本快照/冻结值。"""
    built, texts = _build_and_render()

    for label in ("付款条件", "交付条件", "贸易条款", "客户", "联系人"):
        assert label in texts, f"条款/抬头缺「{label}」：{texts}"
    assert "FOB 宁波" in texts, f"交付条件取版本快照：{texts}"
    assert "虚构客户" in texts and "王经理" in texts, f"抬头取冻结值：{texts}"
    # 付款条件在版本上是 None（未留存）：必须明说，不能空白、更不能拿当前资料补
    assert NOT_RETAINED in texts, f"版本未记录付款条件时必须写「{NOT_RETAINED}」：{texts}"

    for label in ("小计", "运费", "整单优惠", "合计"):
        assert label in texts, f"金额区缺「{label}」：{texts}"
    assert "420.00" in texts and "480.00" in texts, f"金额必须取版本快照：{texts}"
    # 表内明细费用合计对得上版本：小计 + 费用 + 优惠 = 合计
    assert built["subtotal_amount"] + built["charge_amount"] + built["discount_amount"] == (
        built["total_amount"]
    )
    assert built["frozen"]["terms"]["currency"] == "USD"
    assert str(built["frozen"]["items"][0]["unit"]) == "套"
    # 数量按版本行原值带出（Numeric(16,3) 读出来就是 3.500），不做四舍五入
    assert str(built["frozen"]["items"][0]["quantity"]) in ("3.5", "3.500")


def test_legacy_version_unit_is_pending_not_current_sku_unit():
    """旧版本行没有单位快照 → 印"待核实"，**不回查当前 SKU 单位**。"""
    built, texts = _build_and_render(unit_snapshot=None)

    assert UNIT_PENDING in texts, f"缺单位快照时必须写「{UNIT_PENDING}」：{texts}"
    assert "件" not in texts, "绝不能拿当前 SKU 的默认单位（件）冒充历史单位"
    assert any(gap["field"].endswith(".unit") for gap in built["frozen"]["gaps"])


def test_legacy_file_without_frozen_block_is_marked_not_retained():
    """整块冻结信息缺失的历史文件（早于本次修复生成）：显示"未留存"。"""
    legacy = {
        "company_name": "本公司",
        "title": "对客报价单-虚构客户",
        "doc_no": "BJ202610060001",
        "version": 1,
        "created_date": "2026-10-06",
        "status_label": "有效",
        "customer_name": "虚构客户",
        "items": [
            {
                "name": "不锈钢法兰",
                "spec": "DN50",
                "quantity": "5",
                "unit_price": "100",
                "amount": "500",
                "remark": "",
            }
        ],
        "sections": [],
        "charges": [],
        "subtotal_amount": "500",
        "total_amount": "500",
        # 注意：没有 frozen、没有 contact_name —— 就是修复前生成的快照
    }
    texts = _texts(_xlsx(legacy))

    assert "币种" in texts and NOT_RETAINED in texts, f"历史文件必须明说未留存：{texts}"
    assert "USD" not in texts and "CNY" not in texts, "不得凭空补一个币种"
    assert UNIT_PENDING in texts, "历史明细没有单位就写待核实"


def test_doc_pdf_data_carries_frozen_block():
    """下载渲染只认快照：冻结块要原样传给渲染器（否则 Excel 又变回缺币种）。"""
    quote, version, item = _records()
    built = _build(_session(quote, version, item, _charges()))
    snapshot = {key: value for key, value in built.items() if key != "customer"}
    snapshot["customer_name"] = "虚构客户"
    doc = SimpleNamespace(
        doc_no="BJ202610060002",
        doc_type="quote_sheet",
        title="对客报价单-虚构客户",
        version=1,
        created_at=None,
        status="active",
        source_type="quote",
        source_no="Q-CHK-0001",
        source_version=2,
        template_version=1,
        content_sha256="x",
        input_snapshot=snapshot,
        # 第八批 §8.10：`doc_pdf_data` 会把存档信息一并交给渲染器（原件是哪一份、
        # 用哪版渲染器出的）。替身也要有这几列，否则测的就不是真实单据的形状了。
        file_id=None,
        file_sha256=None,
        file_size=None,
        renderer_version=None,
    )
    settings_session = SimpleNamespace(
        execute=AsyncMock(
            return_value=Mock(scalar_one_or_none=Mock(return_value=None))
        )
    )
    data = asyncio.run(doc_pdf_data(settings_session, doc))
    assert data["frozen"]["terms"]["currency"] == "USD"
    assert data["items"][0]["unit"] == "套"
    assert data["contact_name"] == "王经理"
    assert data["owner_name"] == "李销售"
